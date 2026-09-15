"""Render a built report to the formats people actually file and forward.

Every exporter works from one :class:`~pwms.reporting.builder.ReportData`, so a
figure is identical whether it is read on screen, opened in Excel, printed to
PDF or saved as a standalone page:

* ``xlsx`` — a multi-sheet workbook (summary, register, breakdowns, activity,
  referrals) via ``openpyxl``;
* ``csv`` — the flat register, for pipelines and quick imports;
* ``html`` — a self-contained document (inline CSS, no assets to ship);
* ``pdf`` — the same document through WeasyPrint, so the two never drift.

Content builders return ``(filename, bytes, media type)`` so the same bytes can
become a download *or* an email attachment. The heavier libraries are imported
inside their builder rather than at module import: a page view that only renders
HTML should not pay WeasyPrint's or openpyxl's import cost.
"""

from __future__ import annotations

import base64
import datetime as dt
import io

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.text import slugify

from ..models.reports import ATTACHMENT_FORMAT_CHOICES
from .builder import WORKFLOW_COLUMNS

#: Formats offered as an email attachment, with their menu labels. Defined on
#: ``ReportShare`` (which stores one for its scheduled sends) and re-exported
#: here so the reports page and the share record never drift apart.
ATTACHMENT_CHOICES = tuple(ATTACHMENT_FORMAT_CHOICES)

#: Export formats offered by the reports page, in menu order.
EXPORT_FORMATS = tuple(value for value, _label in ATTACHMENT_CHOICES)

#: Media types per format.
CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "html": "text/html; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
}

EXPORT_TEMPLATE = "pwms/pdf/report_export.html"

#: Branding embedded into every export, resolved from the static files on disk.
LOGO_STATIC_PATH = "images/parliament-logo.png"


def suggested_filename(report, extension):
    """A stable, descriptive download name (``overview-report-2026-09-15.xlsx``)."""
    stem = slugify(report.title) or "report"
    stamp = timezone.localtime(report.generated_on).strftime("%Y-%m-%d")
    return f"{stem}-{stamp}.{extension}"


def logo_data_uri():
    """
    The Parliament logo as a ``data:`` URI, or "" when it cannot be read.

    Embedded rather than linked so an export stands alone: WeasyPrint never
    makes a network round trip for it, and the standalone HTML carries its own
    branding.
    """
    from django.contrib.staticfiles import finders

    path = finders.find(LOGO_STATIC_PATH)
    if not path:
        return ""
    try:
        with open(path, "rb") as handle:
            encoded = base64.b64encode(handle.read()).decode("ascii")
    except OSError:
        return ""
    return f"data:image/png;base64,{encoded}"


def build_export_context(report):
    """Context shared by the HTML and PDF renderers (the print document)."""
    return {
        "report": report,
        "document_title": report.title,
        "generated_on": report.generated_on,
        "generated_by": report.user,
        "generated_by_label": report.generated_by_label,
        "filter_summary": report.filter_summary,
        "summary": report.summary,
        "rows": report.rows,
        "by_type": report.by_type,
        "by_state": report.by_state,
        "by_priority": report.by_priority,
        "by_public_status": report.by_public_status,
        "by_owner": report.by_owner,
        "by_group": report.by_group,
        "activity": report.activity,
        "referrals": report.referrals,
        "total_count": report.total,
        "logo_url": logo_data_uri(),
    }


def render_report_html(report):
    """The standalone report document as a string.

    Rendered without a request on purpose: the document needs none of the site
    chrome context processors (navigation, alerts) and must render identically
    from a share or a background render with no request at all.
    """
    return render_to_string(EXPORT_TEMPLATE, build_export_context(report))


# -- content builders -------------------------------------------------------
def export_content(report, export_format):
    """``(filename, content, media type)`` for one format — bytes, always."""
    export_format = (export_format or "xlsx").lower()
    if export_format == "xlsx":
        return (
            suggested_filename(report, "xlsx"),
            _xlsx_bytes(report),
            CONTENT_TYPES["xlsx"],
        )
    if export_format == "pdf":
        return (
            suggested_filename(report, "pdf"),
            _pdf_bytes(report),
            CONTENT_TYPES["pdf"],
        )
    if export_format == "html":
        return (
            suggested_filename(report, "html"),
            render_report_html(report).encode("utf-8"),
            CONTENT_TYPES["html"],
        )
    if export_format == "csv":
        return (
            suggested_filename(report, "csv"),
            _csv_text(report).encode("utf-8"),
            CONTENT_TYPES["csv"],
        )
    raise ValueError(f"Unknown export format: {export_format!r}")


def response_for(report, export_format):
    """An attachment response for ``export_format``."""
    export_format = (export_format or "xlsx").lower()
    if export_format not in EXPORT_FORMATS:
        raise ValueError(f"Unknown export format: {export_format!r}")
    filename, content, media_type = export_content(report, export_format)
    response = HttpResponse(content, content_type=media_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_bytes(report):
    # Imported here: WeasyPrint drags in cairo/pango and is not needed to render
    # the page itself.
    from weasyprint import HTML

    # No base_url: the document is self-contained (inline CSS, embedded logo),
    # so WeasyPrint never needs to fetch anything.
    document = render_report_html(report)
    return HTML(string=document).write_pdf()


def _csv_text(report):
    import csv
    from io import StringIO

    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow([header for header, _key in WORKFLOW_COLUMNS])
    for row in report.rows:
        writer.writerow([_cell(row.get(key)) for _header, key in WORKFLOW_COLUMNS])
    return buffer.getvalue()


def _xlsx_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="275937", end_color="275937", fill_type="solid"
    )

    def add_sheet(title, sheet_headers, sheet_rows, first=False):
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = title
        sheet.append(list(sheet_headers))
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for sheet_row in sheet_rows:
            sheet.append([_cell(value) for value in sheet_row])
        for index, header in enumerate(sheet_headers, 1):
            letter = get_column_letter(index)
            longest = max(
                (
                    len(str(sheet.cell(row=row, column=index).value or ""))
                    for row in range(1, sheet.max_row + 1)
                ),
                default=len(str(header)),
            )
            sheet.column_dimensions[letter].width = min(max(longest + 4, 12), 60)
        return sheet

    summary = report.summary
    summary_rows = [
        ["Report", report.title],
        ["Generated on", _cell(report.generated_on)],
        ["Generated by", report.generated_by_label],
        ["Total workflows", summary["total"]],
        ["Active", summary["active"]],
        ["Completed", summary["completed"]],
        ["Overdue", summary["overdue"]],
        [f"Due within {summary['due_soon_days']} days", summary["due_soon"]],
        ["Unassigned (open)", summary["unassigned"]],
        ["High or urgent priority", summary["high_priority"]],
        ["Referrals", summary["referrals_total"]],
        ["Referrals open", summary["referrals_open"]],
        ["Referrals overdue", summary["referrals_overdue"]],
    ]
    for label in report.filter_summary:
        summary_rows.append(["Filter", label])
    add_sheet("Summary", ["Metric", "Value"], summary_rows, first=True)

    add_sheet(
        "Workflows",
        [header for header, _key in WORKFLOW_COLUMNS],
        [[row.get(key) for _header, key in WORKFLOW_COLUMNS] for row in report.rows],
    )

    _breakdown_sheet(add_sheet, "By Type", report.by_type)
    _breakdown_sheet(add_sheet, "By State", report.by_state)
    _breakdown_sheet(add_sheet, "By Public Status", report.by_public_status)
    _breakdown_sheet(add_sheet, "By Priority", report.by_priority)
    _breakdown_sheet(add_sheet, "By Owner", report.by_owner)

    add_sheet(
        "Activity",
        ["When", "Kind", "Workflow", "From", "To", "Detail", "Actor", "Notes"],
        [
            [
                row["timestamp"],
                row["label"],
                row["workflow"].title,
                row["from_state"],
                row["to_state"],
                row["detail"],
                row["actor"],
                row["notes"],
            ]
            for row in report.activity
        ],
    )

    add_sheet(
        "Referrals",
        [
            "Workflow",
            "Referred to",
            "Referred by",
            "Status",
            "Raised",
            "Due",
            "Overdue",
        ],
        [
            [
                row["workflow"].title,
                row["referred_to"],
                row["referred_by"],
                row["status"],
                row["referred_at"],
                row["due_date"],
                "Yes" if row["is_overdue"] else "No",
            ]
            for row in report.referrals
        ],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _breakdown_sheet(add_sheet, title, breakdown):
    add_sheet(
        title,
        ["Label", "Count", "Share %"],
        [[item["label"], item["count"], item["percent"]] for item in breakdown],
    )


def _cell(value):
    """Make a Python value safe for Excel/CSV (Excel rejects aware datetimes)."""
    if isinstance(value, dt.datetime):
        if timezone.is_aware(value):
            return timezone.localtime(value).replace(tzinfo=None)
        return value
    return "" if value is None else value
