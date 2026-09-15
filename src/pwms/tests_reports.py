"""Reporting: the filter set, the builder, the exports and sharing.

Reporting has six moving parts and each gets its own tests here:

* :class:`~pwms.reporting.builder.ReportFilters` — parsing a request or a stored
  share back into a validated filter set;
* :func:`~pwms.reporting.builder.build_report` — the permission-scoped figures
  (the same chain the lists use, so a report can never out-see its reader);
* :mod:`pwms.reporting.exports` — xlsx / csv / html / pdf;
* :mod:`pwms.reporting.instruments` — one instrument as a formal document;
* :mod:`pwms.reporting.sharing` and the share views — minting a link, emailing
  it, and reopening the *creator's* figures read-only;
* scheduled delivery — which shares repeat, what a send does to the row, the
  management command, and the queued background task.

``build_report`` filters in Python over the accessible set, so most assertions
are about the resulting ``ReportData`` rather than a queryset.
"""

import datetime as dt
from io import BytesIO, StringIO

from background_task.models import Task
from background_task.tasks import tasks as bg_tasks
from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.core import mail
from django.core.management import call_command
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import (
    Attachment,
    Bill,
    BillVersion,
    DelegationParticipant,
    DelegationReport,
    DelegationReportUpdate,
    EventType,
    Group,
    InternationalAgreement,
    InternationalResolution,
    ReportShare,
    State,
    Transition,
    WorkflowEvent,
    WorkflowReferral,
    WorkflowType,
)
from .models.reports import next_occurrence
from .reporting import (
    INSTRUMENT_FORMATS,
    ReportFilters,
    build_instrument_document,
    build_report,
    create_share,
    deliver_scheduled_share,
    due_scheduled_shares,
    export_content,
    instrument_content,
    report_for_share,
    resolve_share,
    send_scheduled_shares,
    suggested_filename,
    suggested_instrument_filename,
)
from .tasks import queue_scheduled_report_shares

User = get_user_model()


class ReportFixtureMixin:
    """A small, fully-known register: two instances the owner can see, one not."""

    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user(
            username="report-owner", email="owner@example.com", password="pw"
        )
        cls.other = User.objects.create_user(
            username="report-other", email="other@example.com", password="pw"
        )
        cls.superuser = User.objects.create_superuser(
            username="report-root", email="root@example.com", password="pw"
        )
        cls.committee = Group.objects.create(name="Reports Committee")

        cls.report_type = WorkflowType.objects.create(name="Reports Test Type")
        cls.draft = State.objects.create(
            workflow_type=cls.report_type,
            name="Drafting",
            public_name="in_progress",
            is_initial=True,
        )
        cls.closed = State.objects.create(
            workflow_type=cls.report_type,
            name="Closed",
            public_name="closed",
            is_terminal=True,
        )
        cls.submit = Transition.objects.create(
            workflow_type=cls.report_type,
            name="Submit",
            from_state=cls.draft,
            to_state=cls.closed,
            requires_comment=False,
        )

        cls.resolution_type = WorkflowType.objects.create(
            name="Reports Resolution Type"
        )
        cls.res_draft = State.objects.create(
            workflow_type=cls.resolution_type,
            name="Captured",
            public_name="new",
            is_initial=True,
        )
        cls.res_closed = State.objects.create(
            workflow_type=cls.resolution_type,
            name="Implemented",
            public_name="implemented",
            is_terminal=True,
        )
        cls.res_implement = Transition.objects.create(
            workflow_type=cls.resolution_type,
            name="Implement",
            from_state=cls.res_draft,
            to_state=cls.res_closed,
            requires_comment=False,
        )

        cls.overdue = DelegationReport.objects.create(
            workflow_type=cls.report_type,
            current_state=cls.draft,
            owner=cls.owner,
            title="Overdue report",
            engagement_name="Engagement A",
            priority="high",
            deadline=timezone.now() - dt.timedelta(days=2),
        )
        cls.resolution = InternationalResolution.objects.create(
            workflow_type=cls.resolution_type,
            current_state=cls.res_draft,
            owner=cls.owner,
            resolution_number="RR-1",
            title="Resolution one",
            priority="urgent",
        )
        # Something the owner may not see: a bill owned by somebody else.
        cls.hidden_bill = Bill.objects.create(
            workflow_type=cls.report_type,
            current_state=cls.draft,
            owner=cls.other,
            bill_number="B 999-2026",
            title="Hidden bill",
        )

        # Activity: one state change on the resolution, one domain event on the
        # report, and one open referral on the report.
        cls.resolution.perform_transition(
            cls.res_implement, actor=cls.owner, comment=""
        )
        event_type = EventType.objects.get_or_create(name="Reports test event")[0]
        WorkflowEvent.objects.create(
            content_type=ContentType.objects.get_for_model(DelegationReport),
            object_id=cls.overdue.pk,
            event_type=event_type,
            actor=cls.owner,
            notes="Recorded for the report test.",
        )
        cls.referral = WorkflowReferral.objects.create(
            content_type=ContentType.objects.get_for_model(DelegationReport),
            object_id=cls.overdue.pk,
            referred_to=cls.committee,
            referred_by=cls.owner,
            due_date=timezone.now() + dt.timedelta(days=3),
        )

    def report(self, **filters):
        return build_report(self.owner, ReportFilters(**filters))


class ReportFilterTests(TestCase):
    """Parsing a request (or a stored share) into a validated filter set."""

    def test_defaults_when_nothing_is_given(self):
        filters = ReportFilters.from_mapping({})
        self.assertEqual(filters.report_type, "overview")
        self.assertEqual(filters.sort, "-created_at")
        self.assertFalse(filters.is_filtered)

    def test_values_are_validated_against_the_choices(self):
        filters = ReportFilters.from_mapping(
            {
                "report_type": "nonsense",
                "priority": "urgent",
                "type": "7",
                "overdue_only": "1",
                "period": "fortnightly",
            }
        )
        # Unknown report type and period fall back to their defaults...
        self.assertEqual(filters.report_type, "overview")
        self.assertEqual(filters.period, "")
        # ...while valid values are kept.
        self.assertEqual(filters.priority, "urgent")
        self.assertEqual(filters.type_id, 7)
        self.assertTrue(filters.overdue_only)
        self.assertTrue(filters.is_filtered)

    def test_non_numeric_ids_are_dropped_rather_than_raising(self):
        filters = ReportFilters.from_mapping(
            {"type": "abc", "owner": "", "state": "3.5"}
        )
        self.assertIsNone(filters.type_id)
        self.assertIsNone(filters.owner_id)
        self.assertIsNone(filters.state_id)

    def test_a_filter_set_round_trips_through_a_dict(self):
        original = ReportFilters(
            report_type="activity",
            q="budget",
            type_id=4,
            priority="high",
            overdue_only=True,
            date_from="2026-01-01",
            date_to="2026-03-31",
            sort="deadline",
        )
        self.assertEqual(ReportFilters.from_dict(original.to_dict()), original)

    def test_period_and_explicit_dates_bound_the_window(self):
        now = timezone.now()
        weekly = ReportFilters(period="weekly")
        start, end = weekly.window(now)
        self.assertEqual(start, now - dt.timedelta(days=7))
        self.assertIsNone(end)

        explicit = ReportFilters(date_from="2026-01-01", date_to="2026-01-31")
        start, end = explicit.window(now)
        self.assertEqual(start.date(), dt.date(2026, 1, 1))
        self.assertEqual(end.date(), dt.date(2026, 1, 31))
        # An explicit "from" beats the period preset.
        combined = ReportFilters(period="yearly", date_from="2026-02-01")
        start, _end = combined.window(now)
        self.assertEqual(start.date(), dt.date(2026, 2, 1))

    def test_bad_dates_are_ignored(self):
        filters = ReportFilters(date_from="not-a-date", date_to="")
        start, end = filters.window(timezone.now())
        self.assertIsNone(start)
        self.assertIsNone(end)


class ReportBuilderTests(ReportFixtureMixin, TestCase):
    """The figures: permission scoping, filters, summary, breakdowns, activity."""

    def test_only_instances_the_reader_may_view_are_counted(self):
        report = self.report()
        self.assertEqual(report.total, 2)
        titles = {row["title"] for row in report.rows}
        self.assertEqual(titles, {"Overdue report", "Resolution one"})

    def test_a_user_only_sees_the_work_that_is_theirs(self):
        report = build_report(self.other, ReportFilters())
        # ``other`` owns the hidden bill, so that one is theirs to see; the
        # owner's two instances are not.
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Hidden bill")

    def test_a_superuser_sees_the_whole_register(self):
        report = build_report(self.superuser, ReportFilters())
        self.assertEqual(report.total, 3)

    def test_type_filter_narrows_to_one_model(self):
        report = self.report(type_id=self.report_type.pk)
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Overdue report")

    def test_state_filter_narrows_by_current_state(self):
        report = self.report(state_id=self.res_closed.pk)
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Resolution one")

    def test_priority_filter(self):
        report = self.report(priority="urgent")
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Resolution one")

    def test_search_matches_title_reference_and_owner(self):
        self.assertEqual(self.report(q="resolution").total, 1)
        self.assertEqual(self.report(q="engAGement a").total, 1)
        self.assertEqual(self.report(q="report-owner").total, 2)
        self.assertEqual(self.report(q="nothing here").total, 0)

    def test_overdue_only_filter(self):
        report = self.report(overdue_only=True)
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Overdue report")

    def test_unassigned_only_filter(self):
        report = self.report(unassigned_only=True)
        self.assertEqual(report.total, 2)

    def test_explicit_date_window_excludes_older_work(self):
        future = timezone.localdate() + dt.timedelta(days=30)
        report = self.report(date_from=future.isoformat())
        self.assertEqual(report.total, 0)

    def test_an_unassigned_instance_can_be_pinned_to_its_assignee(self):
        self.overdue.assigned_to = self.other
        self.overdue.save(update_fields=["assigned_to", "updated_at"])
        # The assignee now sees it (assignment confers view)...
        assigned = build_report(self.other, ReportFilters(assigned_to_id=self.other.pk))
        self.assertEqual(assigned.total, 1)
        # ...and it drops out of the unassigned filter.
        self.assertEqual(self.report(unassigned_only=True).total, 1)

    def test_summary_counts_the_registers_shape(self):
        summary = self.report().summary
        self.assertEqual(summary["total"], 2)
        self.assertEqual(summary["active"], 1)
        self.assertEqual(summary["completed"], 1)
        self.assertEqual(summary["overdue"], 1)
        # Only open work counts as unassigned: the resolution is closed.
        self.assertEqual(summary["unassigned"], 1)
        self.assertEqual(summary["high_priority"], 2)
        self.assertEqual(summary["due_soon"], 0)
        self.assertEqual(summary["referrals_total"], 1)
        self.assertEqual(summary["referrals_open"], 1)

    def test_breakdowns_carry_counts_and_shares(self):
        report = self.report()
        # Breakdowns key on the workflow *type* the instance follows.
        self.assertEqual(
            {item["label"]: item["count"] for item in report.by_type},
            {"Reports Test Type": 1, "Reports Resolution Type": 1},
        )
        self.assertEqual(
            {item["label"]: item["count"] for item in report.by_public_status},
            {"In Progress": 1, "Implemented": 1},
        )
        # The owner holds both, so they get one row carrying both.
        self.assertEqual(
            report.by_owner, [{"label": "report-owner", "count": 2, "percent": 100}]
        )

    def test_priority_breakdown_is_ordered_by_severity(self):
        labels = [item["label"] for item in self.report().by_priority]
        self.assertEqual(labels, ["Urgent", "High"])

    def test_activity_merges_state_changes_and_domain_events(self):
        activity = self.report().activity
        self.assertEqual({row["kind"] for row in activity}, {"transition", "event"})
        # Newest first and attached to the filtered workflows.
        timestamps = [row["timestamp"] for row in activity]
        self.assertEqual(timestamps, sorted(timestamps, reverse=True))
        self.assertEqual(
            {row["workflow"].pk for row in activity},
            {self.overdue.pk, self.resolution.pk},
        )

    def test_referrals_are_collected_for_the_filtered_set(self):
        referrals = self.report().referrals
        self.assertEqual(len(referrals), 1)
        self.assertEqual(referrals[0]["referred_to"], self.committee.name)
        self.assertEqual(referrals[0]["status_value"], "open")

    def test_filter_summary_names_each_active_filter(self):
        report = self.report(
            q="report", priority="high", overdue_only=True, type_id=self.report_type.pk
        )
        summary = " | ".join(report.filter_summary)
        self.assertIn('Search: "report"', summary)
        self.assertIn("Priority: High", summary)
        self.assertIn("Overdue only", summary)
        self.assertIn("Type: Reports Test Type", summary)

    def test_options_are_scoped_to_the_accessible_set(self):
        options = self.report().options
        type_labels = {option["label"] for option in options["types"]}
        self.assertEqual(type_labels, {"Reports Test Type", "Reports Resolution Type"})
        # The hidden bill's owner is not offered as a filter the reader cannot use.
        owner_labels = {option["label"] for option in options["owners"]}
        self.assertNotIn("report-other", owner_labels)


class ReportsPageTests(ReportFixtureMixin, TestCase):
    """The builder page, its preview endpoint and their login gate."""

    def test_the_page_requires_a_login(self):
        response = self.client.get(reverse("pwms:reports"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("pwms:login"), response["Location"])

    def test_the_page_renders_the_filters_and_the_preview(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse("pwms:reports"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Run report")
        self.assertContains(response, "Overdue report")
        self.assertNotContains(response, "Hidden bill")

    def test_the_preview_endpoint_renders_just_the_preview(self):
        self.client.force_login(self.owner)
        response = self.client.get(
            reverse("pwms:reports_preview"), {"report_type": "workflows"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="report-preview"')
        self.assertContains(response, "Overdue report")
        # A preview-only response must not carry the page chrome.
        self.assertNotContains(response, "Run report")

    def test_the_page_honours_filters_from_the_query_string(self):
        self.client.force_login(self.owner)
        response = self.client.get(
            reverse("pwms:reports"), {"report_type": "workflows", "overdue_only": "1"}
        )
        self.assertContains(response, "Overdue report")
        self.assertNotContains(response, "Resolution one")


class ReportExportTests(ReportFixtureMixin, TestCase):
    """Each export format carries the same, complete figures."""

    def setUp(self):
        self.client.force_login(self.owner)

    def _export(self, export_format, **params):
        return self.client.get(
            reverse("pwms:reports_export"), {"format": export_format, **params}
        )

    def test_csv_export_lists_every_matching_workflow(self):
        response = self._export("csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])
        body = response.content.decode("utf-8")
        self.assertIn("Reference", body)
        self.assertIn("Overdue report", body)
        self.assertNotIn("Hidden bill", body)

    def test_xlsx_export_is_a_workbook_with_the_expected_sheets(self):
        response = self._export("xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(response.content))
        for sheet in ("Summary", "Workflows", "By Type", "Activity", "Referrals"):
            self.assertIn(sheet, workbook.sheetnames)
        titles = [
            cell.value
            for cell in workbook["Workflows"]["B"][1:]
            if cell.value is not None
        ]
        self.assertIn("Overdue report", titles)

    def test_html_export_is_a_self_contained_document(self):
        response = self._export("html")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response["Content-Type"])
        body = response.content.decode("utf-8")
        self.assertIn("<!DOCTYPE html>", body)
        self.assertIn("Overdue report", body)
        # Branding is embedded, so the file has no external asset to fetch.
        self.assertIn("data:image/png;base64,", body)

    def test_pdf_export_starts_with_a_pdf_header(self):
        response = self._export("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_an_unknown_format_is_a_bad_request(self):
        self.assertEqual(self._export("docx").status_code, 400)

    def test_exports_are_scoped_to_the_reader(self):
        self.client.force_login(self.other)
        body = self._export("csv").content.decode("utf-8")
        self.assertIn("Hidden bill", body)
        self.assertNotIn("Overdue report", body)

    def test_exports_require_a_login(self):
        self.client.logout()
        self.assertEqual(self._export("csv").status_code, 302)

    def test_suggested_filename_carries_the_title_and_date(self):
        report = build_report(self.owner, ReportFilters())
        name = suggested_filename(report, "xlsx")
        self.assertTrue(name.startswith("overview-"))
        self.assertTrue(name.endswith(".xlsx"))

    def test_export_content_is_bytes_for_every_format(self):
        report = build_report(self.owner, ReportFilters())
        for export_format in ("xlsx", "csv", "html", "pdf"):
            name, content, media = export_content(report, export_format)
            self.assertIsInstance(content, bytes)
            self.assertTrue(media)
            self.assertTrue(name.endswith(export_format))


class ReportShareTests(ReportFixtureMixin, TestCase):
    """Minting a link, emailing it, and reopening the creator's figures."""

    def setUp(self):
        self.client.force_login(self.owner)

    def test_a_link_share_returns_a_url_without_sending_mail(self):
        response = self.client.post(
            reverse("pwms:reports_share"),
            {"report_type": "workflows", "title": "Oversight pack"},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["success"])
        self.assertEqual(payload["sent"], 0)

        share = ReportShare.objects.get()
        self.assertEqual(share.created_by, self.owner)
        self.assertEqual(share.title, "Oversight pack")
        self.assertIn(share.token, payload["link"])
        self.assertEqual(share.filters["report_type"], "workflows")
        self.assertEqual(mail.outbox, [])

    def test_an_email_share_sends_the_link_and_an_attachment(self):
        response = self.client.post(
            reverse("pwms:reports_share"),
            {
                "send_email": "on",
                "recipients": "one@example.com, two@example.com",
                "message": "Have a look.",
                "attach_format": "csv",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["sent"], 2)

        share = ReportShare.objects.get()
        self.assertEqual(share.recipient_list, ["one@example.com", "two@example.com"])

        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertEqual(message.to, ["one@example.com", "two@example.com"])
        self.assertIn(share.token, message.body)
        self.assertIn("Have a look.", message.body)
        self.assertEqual(len(message.attachments), 1)
        self.assertTrue(message.attachments[0][0].endswith(".csv"))

    def test_emailing_without_recipients_is_rejected(self):
        response = self.client.post(reverse("pwms:reports_share"), {"send_email": "on"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("recipients", response.json()["errors"])
        self.assertFalse(ReportShare.objects.exists())

    def test_a_malformed_recipient_is_rejected(self):
        response = self.client.post(
            reverse("pwms:reports_share"),
            {"send_email": "on", "recipients": "not-an-address"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("recipients", response.json()["errors"])

    def test_the_share_flow_requires_a_post(self):
        self.assertEqual(
            self.client.get(reverse("pwms:reports_share")).status_code, 405
        )

    def test_a_shared_link_reopens_the_creators_figures_read_only(self):
        share = ReportShare.objects.create(
            created_by=self.owner,
            title="Shared overview",
            filters=ReportFilters().to_dict(),
        )
        # ``other`` cannot see this work in their own reports...
        self.assertEqual(build_report(self.other, ReportFilters()).total, 1)

        self.client.force_login(self.other)
        response = self.client.get(reverse("pwms:report_shared", args=[share.token]))
        self.assertEqual(response.status_code, 200)
        # ...but the shared page shows exactly what the creator could see.
        self.assertContains(response, "Overdue report")
        self.assertContains(response, "Shared report")
        self.assertNotContains(response, 'id="report-share-modal"')

        share.refresh_from_db()
        self.assertEqual(share.access_count, 1)
        self.assertIsNotNone(share.last_accessed_at)

    def test_exporting_from_a_share_uses_the_creators_figures(self):
        share = ReportShare.objects.create(
            created_by=self.owner, filters=ReportFilters().to_dict()
        )
        self.client.force_login(self.other)
        response = self.client.get(
            reverse("pwms:reports_export"),
            {"share": share.token, "format": "csv"},
        )
        self.assertEqual(response.status_code, 200)
        body = response.content.decode("utf-8")
        self.assertIn("Overdue report", body)
        self.assertNotIn("Hidden bill", body)

    def test_an_unknown_token_is_not_found(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse("pwms:report_shared", args=["nope"]))
        self.assertEqual(response.status_code, 404)

    def test_a_revoked_link_is_gone(self):
        share = ReportShare.objects.create(
            created_by=self.owner, filters=ReportFilters().to_dict()
        )
        share.revoke()
        self.client.force_login(self.other)
        response = self.client.get(reverse("pwms:report_shared", args=[share.token]))
        self.assertEqual(response.status_code, 410)
        self.assertContains(response, "no longer active", status_code=410)

    def test_an_expired_link_is_gone(self):
        share = ReportShare.objects.create(
            created_by=self.owner,
            filters=ReportFilters().to_dict(),
            expires_at=timezone.now() - dt.timedelta(minutes=1),
        )
        self.client.force_login(self.other)
        response = self.client.get(reverse("pwms:report_shared", args=[share.token]))
        self.assertEqual(response.status_code, 410)

    def test_a_share_whose_creator_is_gone_is_unreadable(self):
        share = ReportShare.objects.create(
            created_by=self.owner, filters=ReportFilters().to_dict()
        )
        ReportShare.objects.filter(pk=share.pk).update(created_by=None)

        self.client.force_login(self.other)
        response = self.client.get(reverse("pwms:report_shared", args=[share.token]))
        self.assertEqual(response.status_code, 410)
        self.assertContains(response, "no longer available", status_code=410)
        self.assertIsNone(report_for_share(resolve_share(share.token)))

    def test_resolve_share_finds_the_row_and_report_for_share_rebuilds_it(self):
        share = ReportShare.objects.create(
            created_by=self.owner, filters=ReportFilters(overdue_only=True).to_dict()
        )
        resolved = resolve_share(share.token)
        self.assertEqual(resolved.pk, share.pk)

        report = report_for_share(resolved)
        self.assertEqual(report.total, 1)
        self.assertEqual(report.rows[0]["title"], "Overdue report")

        self.assertIsNone(resolve_share("no-such-token"))


class ReportShareModelTests(TestCase):
    """The share row's own state helpers."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username="share-owner", password="pw")

    def test_recipients_are_split_and_trimmed(self):
        share = ReportShare.objects.create(
            created_by=self.user,
            recipients=" a@example.com ;b@example.com, ,c@example.com ",
        )
        self.assertEqual(
            share.recipient_list,
            ["a@example.com", "b@example.com", "c@example.com"],
        )

    def test_tokens_are_unique_and_non_trivial(self):
        first = ReportShare.objects.create(created_by=self.user)
        second = ReportShare.objects.create(created_by=self.user)
        self.assertNotEqual(first.token, second.token)
        self.assertGreaterEqual(len(first.token), 32)

    def test_active_until_revoked_or_expired(self):
        live = ReportShare.objects.create(created_by=self.user)
        self.assertTrue(live.is_active)
        self.assertFalse(live.is_expired)

        expired = ReportShare.objects.create(
            created_by=self.user, expires_at=timezone.now() - dt.timedelta(seconds=1)
        )
        self.assertFalse(expired.is_active)
        self.assertTrue(expired.is_expired)

        live.revoke()
        self.assertFalse(live.is_active)
        # Revoking twice is harmless.
        live.revoke()
        self.assertIsNotNone(live.revoked_at)

    def test_record_access_counts_each_open(self):
        share = ReportShare.objects.create(created_by=self.user)
        share.record_access()
        share.record_access()
        share.refresh_from_db()
        self.assertEqual(share.access_count, 2)
        self.assertIsNotNone(share.last_accessed_at)


class ReportShareAdminTests(TestCase):
    """The admin's revoke action is the one write offered for a share row."""

    @classmethod
    def setUpTestData(cls):
        cls.staff = User.objects.create_superuser(
            username="share-admin", email="share-admin@example.com", password="pw"
        )
        cls.owner = User.objects.create_user(username="share-maker", password="pw")

    def test_the_revoke_action_cuts_the_selected_links_off(self):
        first = ReportShare.objects.create(created_by=self.owner, title="First")
        second = ReportShare.objects.create(created_by=self.owner, title="Second")
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("admin:pwms_reportshare_changelist"),
            {"action": "revoke_shares", "_selected_action": [first.pk, second.pk]},
        )
        self.assertEqual(response.status_code, 302)

        first.refresh_from_db()
        second.refresh_from_db()
        self.assertFalse(first.is_active)
        self.assertFalse(second.is_active)


class ReportingSmokeTests(TestCase):
    """The app boots with the reporting package importable and tables present."""

    def test_report_share_table_exists(self):
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT 1 FROM pg_catalog.pg_tables WHERE tablename = %s",
                ["pwms_reportshare"],
            )
            self.assertIsNotNone(cursor.fetchone())


# ---------------------------------------------------------------------------
# Scheduled delivery
# ---------------------------------------------------------------------------


def _labels(section):
    """The labels of one fact section, in order."""
    return [field["label"] for field in section["fields"]]


def _section(document, title):
    """The fact section titled ``title``."""
    for section in document.sections:
        if section["title"] == title:
            return section
    raise AssertionError(f"no section titled {title!r}")


def _table(document, prefix):
    """The table whose title starts with ``prefix``."""
    for table in document.tables:
        if table["title"].startswith(prefix):
            return table
    raise AssertionError(f"no table starting with {prefix!r}")


def _rows(table):
    """A table's cells as plain strings."""
    return [[cell["text"] for cell in row] for row in table["rows"]]


class SchedulingHelperTests(TestCase):
    """The calendar arithmetic behind a share's next send."""

    def setUp(self):
        self.start = dt.datetime(2026, 1, 31, 8, 0, tzinfo=dt.UTC)

    def test_daily_and_weekly_are_fixed_intervals(self):
        self.assertEqual(
            next_occurrence(self.start, "daily"), self.start + dt.timedelta(days=1)
        )
        self.assertEqual(
            next_occurrence(self.start, "weekly"), self.start + dt.timedelta(days=7)
        )

    def test_an_unscheduled_share_has_no_next_occurrence(self):
        self.assertIsNone(next_occurrence(self.start, "none"))
        self.assertIsNone(next_occurrence(self.start, ""))

    def test_monthly_clamps_to_the_target_months_length(self):
        # 31 January -> 28 February (2026 is not a leap year), keeping the time.
        result = next_occurrence(self.start, "monthly")
        self.assertEqual((result.year, result.month, result.day), (2026, 2, 28))
        self.assertEqual(result.hour, 8)

    def test_monthly_rolls_the_year_over(self):
        december = dt.datetime(2026, 12, 15, 8, 0, tzinfo=dt.UTC)
        result = next_occurrence(december, "monthly")
        self.assertEqual((result.year, result.month, result.day), (2027, 1, 15))


class ScheduledShareFixtureMixin(ReportFixtureMixin):
    """A helper for minting the kind of share the delivery job looks for."""

    def due_share(self, **overrides):
        values = {
            "created_by": self.owner,
            "title": "Weekly oversight",
            "recipients": "one@example.com, two@example.com",
            "schedule": "weekly",
            "schedule_format": "csv",
            "next_send_at": timezone.now() - dt.timedelta(minutes=5),
            "filters": ReportFilters().to_dict(),
        }
        values.update(overrides)
        return ReportShare.objects.create(**values)


class ScheduledShareTests(ScheduledShareFixtureMixin, TestCase):
    """Which shares are due, and what a send does to the row."""

    def test_only_active_due_repeating_shares_are_queued(self):
        due = self.due_share()
        self.due_share(next_send_at=timezone.now() + dt.timedelta(days=1))
        self.due_share(schedule="none", next_send_at=None)
        self.due_share(revoked_at=timezone.now())
        self.due_share(expires_at=timezone.now() - dt.timedelta(minutes=1))

        self.assertEqual(list(due_scheduled_shares()), [due])

    def test_a_send_emails_the_link_and_advances_the_schedule(self):
        share = self.due_share()
        before = share.next_send_at

        result = deliver_scheduled_share(share)

        self.assertTrue(result.sent)
        self.assertEqual(result.recipients, 2)
        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertEqual(message.to, ["one@example.com", "two@example.com"])
        self.assertIn(share.token, message.body)
        self.assertEqual(len(message.attachments), 1)

        share.refresh_from_db()
        self.assertEqual(share.send_count, 1)
        self.assertIsNotNone(share.last_sent_at)
        self.assertEqual(share.last_error, "")
        self.assertGreater(share.next_send_at, before)
        self.assertFalse(share.is_due())

    def test_a_share_with_nowhere_to_send_records_why_and_still_advances(self):
        share = self.due_share(recipients="")
        before = share.next_send_at

        result = deliver_scheduled_share(share)

        self.assertFalse(result.sent)
        self.assertEqual(result.recipients, 0)
        self.assertEqual(mail.outbox, [])

        share.refresh_from_db()
        self.assertIn("no recipients", share.last_error)
        self.assertEqual(share.send_count, 1)
        # Advanced, so a broken share retries next interval, not on every run.
        self.assertGreater(share.next_send_at, before)

    def test_a_successful_send_clears_an_earlier_error(self):
        share = self.due_share(last_error="smtp was down")

        deliver_scheduled_share(share)

        share.refresh_from_db()
        self.assertEqual(share.last_error, "")
        self.assertTrue(share.is_scheduled)

    def test_the_batch_summarises_what_it_did(self):
        self.due_share()
        self.due_share(recipients="")

        report = send_scheduled_shares()

        self.assertEqual(
            report["summary"],
            {"considered": 2, "sent": 1, "failed": 1, "recipients": 2},
        )
        self.assertEqual(len(report["results"]), 2)

    def test_create_share_books_the_first_repeat_one_interval_out(self):
        share = create_share(
            user=self.owner,
            filters=ReportFilters(),
            schedule="daily",
            schedule_format="pdf",
        )

        self.assertTrue(share.is_scheduled)
        self.assertEqual(share.schedule_format, "pdf")
        self.assertIsNotNone(share.next_send_at)
        self.assertGreater(share.next_send_at, timezone.now())
        self.assertFalse(share.is_due())

    def test_an_unscheduled_share_never_falls_due(self):
        share = create_share(user=self.owner, filters=ReportFilters())

        self.assertFalse(share.is_scheduled)
        self.assertIsNone(share.next_send_at)
        self.assertFalse(share.is_due())


class ScheduledDeliveryCommandTests(ScheduledShareFixtureMixin, TestCase):
    """The command a cron job - or a worker draining the queue - runs."""

    def test_it_sends_the_due_shares(self):
        share = self.due_share()
        out = StringIO()

        call_command("send_scheduled_report_shares", stdout=out)

        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("1 due, 1 sent, 0 failed", out.getvalue())
        share.refresh_from_db()
        self.assertEqual(share.send_count, 1)

    def test_dry_run_sends_nothing_and_saves_nothing(self):
        share = self.due_share()
        before = share.next_send_at
        out = StringIO()

        call_command("send_scheduled_report_shares", "--dry-run", stdout=out)

        self.assertEqual(mail.outbox, [])
        self.assertIn("DRY RUN", out.getvalue())
        self.assertIn("Would send", out.getvalue())
        share.refresh_from_db()
        self.assertEqual(share.send_count, 0)
        self.assertEqual(share.next_send_at, before)

    def test_queueing_registers_a_repeating_background_task(self):
        out = StringIO()

        call_command("send_scheduled_report_shares", "--queue", stdout=out)

        self.assertIn("Queued", out.getvalue())
        task = Task.objects.get()
        self.assertEqual(task.repeat, Task.HOURLY)
        self.assertEqual(task.task_name, "pwms.tasks.deliver_scheduled_report_shares")
        # Queuing alone must not deliver anything.
        self.assertEqual(mail.outbox, [])


class ScheduledDeliveryTaskTests(ScheduledShareFixtureMixin, TestCase):
    """The background-task handle the worker claims."""

    def test_the_queued_task_delivers_the_due_shares(self):
        share = self.due_share()

        with self.settings(BACKGROUND_TASK_RUN_ASYNC=False):
            task = queue_scheduled_report_shares()
            bg_tasks.run_task(task)

        self.assertEqual(len(mail.outbox), 1)
        share.refresh_from_db()
        self.assertEqual(share.send_count, 1)

    def test_the_queued_task_is_a_repeating_hourly_drain(self):
        with self.settings(BACKGROUND_TASK_RUN_ASYNC=False):
            task = queue_scheduled_report_shares()

        self.assertEqual(task.repeat, Task.HOURLY)
        self.assertEqual(task.task_name, "pwms.tasks.deliver_scheduled_report_shares")


class ReportShareScheduleViewTests(ScheduledShareFixtureMixin, TestCase):
    """Scheduling a share from the reports page."""

    def setUp(self):
        self.client.force_login(self.owner)

    def test_scheduling_emails_now_and_books_the_next_send(self):
        response = self.client.post(
            reverse("pwms:reports_share"),
            {
                "title": "Standing oversight report",
                "schedule": "weekly",
                "send_email": "on",
                "recipients": "one@example.com",
                "attach_format": "csv",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["schedule"], "weekly")
        self.assertIsNotNone(payload["next_send_at"])
        self.assertEqual(payload["sent"], 1)

        share = ReportShare.objects.get()
        self.assertEqual(share.schedule, "weekly")
        self.assertEqual(share.schedule_format, "csv")
        self.assertTrue(share.is_scheduled)
        self.assertEqual(len(mail.outbox), 1)

    def test_scheduling_without_recipients_is_rejected(self):
        response = self.client.post(
            reverse("pwms:reports_share"), {"schedule": "daily"}
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("recipients", response.json()["errors"])
        self.assertFalse(ReportShare.objects.exists())

    def test_scheduling_implies_recipients_without_ticking_the_box(self):
        response = self.client.post(
            reverse("pwms:reports_share"),
            {"schedule": "monthly", "recipients": "one@example.com"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(ReportShare.objects.get().recipient_list, ["one@example.com"])

    def test_a_one_off_link_is_not_scheduled(self):
        response = self.client.post(reverse("pwms:reports_share"), {})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["schedule"], "none")
        share = ReportShare.objects.get()
        self.assertFalse(share.is_scheduled)
        self.assertIsNone(share.next_send_at)


# ---------------------------------------------------------------------------
# Per-instrument formal documents
# ---------------------------------------------------------------------------


class InstrumentFixtureMixin:
    """One of every instrument, carrying the detail a document should print."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username="doc-owner", email="doc-owner@example.com", password="pw"
        )
        cls.other = User.objects.create_user(username="doc-other", password="pw")
        cls.committee = Group.objects.create(name="Document Committee")

        cls.wt = WorkflowType.objects.create(name="Document Test Type")
        cls.draft = State.objects.create(
            workflow_type=cls.wt,
            name="Drafting",
            public_name="in_progress",
            is_initial=True,
        )
        cls.closed = State.objects.create(
            workflow_type=cls.wt,
            name="Closed",
            public_name="closed",
            is_terminal=True,
        )
        cls.submit = Transition.objects.create(
            workflow_type=cls.wt,
            name="Submit",
            from_state=cls.draft,
            to_state=cls.closed,
            requires_comment=False,
        )

        cls.report = DelegationReport.objects.create(
            workflow_type=cls.wt,
            current_state=cls.draft,
            owner=cls.user,
            title="Annual delegation report",
            engagement_name="Inter-Parliamentary Forum",
            engagement_start_date=dt.date(2026, 3, 1),
            engagement_end_date=dt.date(2026, 3, 5),
            notes="Follow up with the secretariat.",
            priority="high",
        )
        DelegationParticipant.objects.create(
            delegation_report=cls.report,
            participant_type=DelegationParticipant.MEMBER,
            title="Hon.",
            first_name="Naledi",
            last_name="Mokoena",
            delegation_role="Leader of the Delegation",
        )
        DelegationReportUpdate.objects.create(
            delegation_report=cls.report,
            update_date=dt.date(2026, 4, 1),
            resulting_state=cls.draft,
            atc_reference="ATC 12-2026",
            atc_publication_date=dt.date(2026, 4, 2),
            atc_page_number="1145",
            notes="Tabled.",
            recorded_by=cls.user,
        )
        Attachment.objects.create(
            content_type=ContentType.objects.get_for_model(DelegationReport),
            object_id=str(cls.report.pk),
            name="Report.pdf",
            drive_id="drive-1",
            item_id="item-1",
            type="report",
            uploaded_by=cls.user,
        )

        cls.resolution = InternationalResolution.objects.create(
            workflow_type=cls.wt,
            current_state=cls.draft,
            owner=cls.user,
            resolution_number="RR-42",
            title="Resolution on climate finance",
            resolution_text="The delegation resolves to pursue climate finance.",
            implementation_progress="Committee briefed.",
            adoption_date=dt.date(2026, 3, 4),
        )
        cls.report.add_sub_workflow(cls.resolution)

        cls.agreement = InternationalAgreement.objects.create(
            workflow_type=cls.wt,
            current_state=cls.draft,
            owner=cls.user,
            title="Extradition treaty",
            agreement_type=InternationalAgreement.SECTION_231_2,
            submitting_department="Department of Justice",
            responsible_minister_name="Minister of Justice",
            atc_tabling_date=dt.date(2026, 2, 10),
            atc_reference="ATC 3-2026",
        )
        cls.agreement.referral_committees.add(cls.committee)

        cls.bill = Bill.objects.create(
            workflow_type=cls.wt,
            current_state=cls.draft,
            owner=cls.user,
            bill_number="B 7-2026",
            title="Climate Finance Bill",
            short_title="Climate Finance",
            bill_type=Bill.SECTION_76,
            house_of_origin=Bill.NA,
            sponsor_name="Minister of Finance",
            introduced_date=dt.date(2026, 5, 6),
            atc_reference="ATC 20-2026",
        )
        BillVersion.objects.create(
            bill=cls.bill,
            version_label="B 7-2026 (1st)",
            version_type=BillVersion.INTRODUCED,
            version_date=dt.date(2026, 5, 6),
            is_current=True,
            recorded_by=cls.user,
            notes="As introduced.",
        )

        # Activity on the report, and a referral, so those tables have rows.
        cls.report.perform_transition(cls.submit, actor=cls.user, comment="")
        WorkflowReferral.objects.create(
            content_type=ContentType.objects.get_for_model(DelegationReport),
            object_id=cls.report.pk,
            referred_to=cls.committee,
            referred_by=cls.user,
        )


class InstrumentDocumentTests(InstrumentFixtureMixin, TestCase):
    """The facts, notes and tables each instrument document is built from."""

    def test_a_delegation_report_gathers_its_engagement_and_lists(self):
        document = build_instrument_document(self.report, self.user)

        self.assertEqual(document.document_title, "Annual delegation report")
        self.assertIn("Document Test Type", document.document_subtitle)
        self.assertIn(self.report.identifier, document.document_subtitle)

        self.assertEqual(
            [section["title"] for section in document.sections][:2],
            ["At a glance", "Engagement"],
        )
        # The report has a child, so completeness is reported too.
        self.assertIn("Sub-instrument completion", _labels(document.sections[0]))
        self.assertIn("Latest ATC reference", _labels(_section(document, "Engagement")))

        for prefix in (
            "Delegation participants",
            "BR03 updates",
            "Related instruments",
            "Referrals",
            "Documents",
        ):
            self.assertTrue(_table(document, prefix)["rows"])
        self.assertTrue(_table(document, "State changes")["rows"])
        self.assertTrue(_table(document, "Domain events")["rows"])
        self.assertIn("Notes / follow-up", [note["title"] for note in document.notes])

    def test_the_participant_and_update_rows_carry_their_values(self):
        document = build_instrument_document(self.report, self.user)

        self.assertEqual(
            _rows(_table(document, "Delegation participants")),
            [
                [
                    "Hon. Naledi Mokoena",
                    "Leader of the Delegation",
                    "Parliamentary Delegation Member",
                    "—",
                ]
            ],
        )
        self.assertEqual(_rows(_table(document, "BR03 updates"))[0][2], "ATC 12-2026")

    def test_a_resolution_document_states_its_text_and_its_parent(self):
        document = build_instrument_document(self.resolution, self.user)

        self.assertEqual(
            _section(document, "Resolution")["fields"][0]["value"], "RR-42"
        )
        self.assertIn("Resolution text", [note["title"] for note in document.notes])
        self.assertIn(
            "Implementation progress", [note["title"] for note in document.notes]
        )
        self.assertEqual(
            _rows(_table(document, "Parent instrument"))[0][1],
            "Annual delegation report",
        )

    def test_an_agreement_document_carries_its_department_and_committees(self):
        document = build_instrument_document(self.agreement, self.user)
        fields = {
            field["label"]: field["value"]
            for field in _section(document, "Agreement")["fields"]
        }

        self.assertEqual(fields["Submitting department"], "Department of Justice")
        self.assertEqual(fields["Agreement type"], "Section 231(2) agreement")
        self.assertEqual(fields["Referral committees"], "Document Committee")

    def test_a_bill_document_lists_its_profile_and_versions(self):
        document = build_instrument_document(self.bill, self.user)

        titles = [section["title"] for section in document.sections]
        self.assertIn("Bill profile", titles)
        self.assertIn("References", titles)
        fields = {
            field["label"]: field["value"]
            for field in _section(document, "Bill profile")["fields"]
        }
        self.assertEqual(fields["Sponsor"], "Minister of Finance")
        self.assertEqual(
            _rows(_table(document, "Bill versions"))[0][0], "B 7-2026 (1st)"
        )

    def test_the_html_document_is_self_contained_for_every_type(self):
        for instance in (self.report, self.resolution, self.agreement, self.bill):
            with self.subTest(kind=type(instance).__name__):
                name, content, media = instrument_content(instance, self.user, "html")
                body = content.decode("utf-8")
                self.assertTrue(name.endswith(".html"))
                self.assertIn("text/html", media)
                self.assertIn("<!DOCTYPE html>", body)
                self.assertIn(instance.title, body)
                # Branding is embedded, so the file has nothing external to fetch.
                self.assertIn("data:image/png;base64,", body)

    def test_the_document_renders_to_pdf(self):
        name, content, media = instrument_content(self.bill, self.user, "pdf")

        self.assertTrue(name.endswith(".pdf"))
        self.assertEqual(media, "application/pdf")
        self.assertTrue(content.startswith(b"%PDF"))

    def test_an_unknown_format_is_rejected(self):
        self.assertEqual(INSTRUMENT_FORMATS, ("pdf", "html"))
        with self.assertRaises(ValueError):
            instrument_content(self.report, self.user, "docx")

    def test_the_filename_carries_the_reference_and_the_type(self):
        name = suggested_instrument_filename(self.bill, "pdf")

        self.assertIn("b-7-2026", name)
        self.assertIn("bill", name)
        self.assertTrue(name.endswith(".pdf"))


class InstrumentDocumentViewTests(InstrumentFixtureMixin, TestCase):
    """Downloading one instrument's formal document."""

    def _url(self, instance):
        return reverse("pwms:workflow_document", args=[instance.public_id])

    def test_the_document_requires_a_login(self):
        response = self.client.get(self._url(self.report))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("pwms:login"), response["Location"])

    def test_a_reader_without_access_is_refused(self):
        self.client.force_login(self.other)

        response = self.client.get(self._url(self.report))

        self.assertEqual(response.status_code, 403)

    def test_an_unknown_instance_is_not_found(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse(
                "pwms:workflow_document",
                args=["00000000-0000-0000-0000-000000000000"],
            )
        )

        self.assertEqual(response.status_code, 404)

    def test_the_document_downloads_as_pdf_by_default(self):
        self.client.force_login(self.user)

        response = self.client.get(self._url(self.resolution))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_the_document_can_be_opened_as_a_page(self):
        self.client.force_login(self.user)

        response = self.client.get(self._url(self.agreement), {"format": "html"})

        self.assertEqual(response.status_code, 200)
        self.assertIn("text/html", response["Content-Type"])
        self.assertIn("Extradition treaty", response.content.decode("utf-8"))

    def test_an_unknown_format_is_a_bad_request(self):
        self.client.force_login(self.user)

        response = self.client.get(self._url(self.report), {"format": "docx"})

        self.assertEqual(response.status_code, 400)

    def test_every_detail_page_offers_the_document(self):
        self.client.force_login(self.user)

        for instance in (self.report, self.resolution, self.agreement, self.bill):
            with self.subTest(kind=type(instance).__name__):
                response = self.client.get(instance.get_absolute_url())
                self.assertContains(response, self._url(instance))
